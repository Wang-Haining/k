"Score human agreement and weighted instrument accuracy from private returned ratings."

import argparse
import ast
import csv
import json
import math
import re
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config import DATA_DIR, RESULTS_DIR, ROOT

CHOICES = next(
    ast.literal_eval(n.value)
    for n in ast.parse((ROOT / "validation" / "sample.py").read_text()).body
    if isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "CHOICES" for t in n.targets)
)
SHEETS = dict(
    zip(
        ("dossier", "dossier_trial", "intent", "boundary", "lineage"),
        ("dossier_awardees", "dossier_trials", "intent", "boundary", "lineage"),
        strict=False,
    )
)
PAIR = ("identity", "role", "scope", "k_relationship")
FIELDS = {
    "dossier": ("led_K_trial", "led_new_trial", "search_done"),
    "dossier_trial": (*PAIR, "start_in_window"),
    "intent": ("intent",),
    "boundary": ("k_relationship",),
    "lineage": ("lineage",),
}
FREE = {"missed_trials_found", "notes", "reason", "evidence", "rater_notes"}
QUAL = {"overall_PI", "responsible_party_PI", "sponsor_investigator"}
DEV = {"extension", "same_approach_new_question"}
NCT = re.compile("\\bNCT\\d{8}\\b", re.IGNORECASE)
KEY_COLUMNS = (
    "module",
    "item",
    "blind_id",
    "case_id",
    "stratum",
    "intent",
    "instrument_state",
    "cell_population",
    "instrument_label",
)
ADJ_COLUMNS = ("module", "item", "field", "rater_a", "rater_b", "final")
REPS = 2000


def require(ok, message):
    if not ok:
        raise ValueError(message)


def text(value):
    return "" if value is None else str(value).strip()


def choices(module, field):
    kind = "yes_no_unclear" if field in FIELDS["dossier"] else field
    if module == "boundary":
        kind = "krel"
    return CHOICES[kind]


def ncts(value):
    return {x.upper() for x in NCT.findall(value)}


def answer(module, field, value):
    if field == "missed_trials_found":
        ids = ncts(value)
        require(not value or ids, f"{module}/{field}: nonempty text has no NCT id: {value!r}")
        require(
            not re.search("\\bNCT(?!\\d{8}\\b)", value, re.IGNORECASE), f"{module}/{field}: malformed NCT id: {value!r}"
        )
        return tuple(sorted(ids))
    require(value in choices(module, field), f"{module}/{field}: expected {choices(module, field)}, got {value!r}")
    return value


def csv_rows(path, columns):
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        require(
            reader.fieldnames and set(columns) <= set(reader.fieldnames),
            f"{path}: missing columns; expected {columns}, got {reader.fieldnames}",
        )
        require(len(reader.fieldnames) == len(set(reader.fieldnames)), f"{path}: duplicate columns")
        rows = list(reader)
    require(rows, f"{path}: expected N>0 rows, got 0")
    for i, row in enumerate(rows, 2):
        require(None not in row and all(v is not None for v in row.values()), f"{path}:{i}: malformed row {row}")
    return rows


def indexed(rows, field, context):
    out = {}
    for row in rows:
        key = row[field]
        require(key and key not in out, f"{context}: empty/duplicate {field}={key!r}")
        out[key] = row
    return out


def workbook(path):
    wb = load_workbook(path, read_only=True, data_only=False)
    out = {}
    try:
        require(all(s in wb.sheetnames for s in SHEETS.values()), f"{path}: missing sheets; got {wb.sheetnames}")
        for module, sheet in SHEETS.items():
            rows = iter(wb[sheet].iter_rows(values_only=True))
            cols = tuple(text(x) for x in next(rows))
            required = set(FIELDS[module]) | {"item"}
            required |= {"trial", "nct_id"} if module == "dossier_trial" else set()
            required |= {"missed_trials_found"} if module == "dossier" else set()
            require(
                required <= set(cols) and len(cols) == len(set(cols)) and all(cols),
                f"{path}/{sheet}: missing/duplicate/empty columns: {cols}",
            )
            for raw in rows:
                if all(v is None for v in raw):
                    continue
                row = dict(zip(cols, map(text, raw), strict=False))
                item = row["trial" if module == "dossier_trial" else "item"]
                require(item and (module, item) not in out, f"{path}/{sheet}: empty/duplicate item {item!r}")
                for field in FIELDS[module]:
                    answer(module, field, row[field])
                if module == "dossier":
                    answer(module, "missed_trials_found", row["missed_trials_found"])
                out[module, item] = row
            require(any(k[0] == module for k in out), f"{path}/{sheet}: expected N>0 rows, got 0")
        require(out, f"{path}: expected ratings, got 0 rows")
    finally:
        wb.close()
    return out


def ratings(base):
    books = [workbook(base / "returned" / f"rating_{r}.xlsx") for r in ("rater_a", "rater_b")]
    a, b = books
    require(a.keys() == b.keys(), f"Rater item sets differ: {a.keys() ^ b.keys()}")
    disagreements = {}
    for key, row in a.items():
        module, item = key
        editable = set(FIELDS[module]) | FREE
        require(
            {k: v for k, v in row.items() if k not in editable}
            == {k: v for k, v in b[key].items() if k not in editable},
            f"{key}: raters have different supplied data",
        )
        for field in (*FIELDS[module], *(["missed_trials_found"] if module == "dossier" else [])):
            if answer(module, field, row[field]) != answer(module, field, b[key][field]):
                disagreements[module, item, field] = (row[field], b[key][field])
    return (a, b, disagreements)


def merge(base):
    _, _, dis = ratings(base)
    path = base / "adjudication.xlsx"
    require(not path.exists(), f"{path}: already exists; preserve completed adjudication before rerunning merge")
    wb = Workbook()
    ws = wb.active
    ws.title = "adjudication"
    ws.append(ADJ_COLUMNS)
    for (module, item, field), values in dis.items():
        ws.append((module, item, field, *values, None))
        row = ws.max_row
        ws.cell(row, 6).fill = PatternFill("solid", fgColor="FFF2CC")
        if field != "missed_trials_found":
            dv = DataValidation(type="list", formula1='"' + ",".join(choices(module, field)) + '"')
            ws.add_data_validation(dv)
            dv.add(f"F{row}")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "D2"
    for col in "ABCDEF":
        ws.column_dimensions[col].width = 30 if col in "ABC" else 65
    wb.save(path)
    return dis


def reference(base, a, dis):
    ref = {k: dict(row) for k, row in a.items()}
    path = base / "returned" / "adjudication.xlsx"
    require(path.exists() or not dis, f"{len(dis)} unresolved disagreements; expected {path}")
    if not path.exists():
        return ref
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        require(wb.sheetnames == ["adjudication"], f"{path}: unexpected sheets {wb.sheetnames}")
        rows = iter(wb.active.iter_rows(values_only=True))
        require(tuple(next(rows)) == ADJ_COLUMNS, f"{path}: expected columns {ADJ_COLUMNS}")
        seen = set()
        for raw in rows:
            if all(v is None for v in raw):
                continue
            module, item, field, av, bv, final = map(text, raw)
            key = (module, item, field)
            require(key in dis and key not in seen, f"{path}: stale/duplicate adjudication {key}")
            require((av, bv) == dis[key], f"{key}: adjudication answers differ from returned ratings")
            require(bool(final), f"{key}: adjudication final is blank; use none for an empty missed-trial set")
            if field == "missed_trials_found" and final.lower() == "none":
                final = ""
            answer(module, field, final)
            ref[module, item][field] = final
            seen.add(key)
        require(seen == dis.keys(), f"Unresolved adjudication cells: {dis.keys() - seen}")
    finally:
        wb.close()
    return ref


def agreement(module, field, pairs, comparison="rater_a_vs_rater_b", total=None):
    n = len(pairs)
    matches = sum((a == b for a, b in pairs))
    ca, cb = (Counter((a for a, _ in pairs)), Counter((b for _, b in pairs)))
    expected = sum(ca[k] * cb[k] for k in ca) / n**2 if n else 1
    return {
        "module": module,
        "field": field,
        "comparison": comparison,
        "n_total": n if total is None else total,
        "n_compared": n,
        "n_excluded": 0 if total is None else total - n,
        "percent_agreement": 100 * matches / n if n else "",
        "kappa": (matches / n - expected) / (1 - expected) if n and expected < 1 else "",
        "status": "ok" if n and expected < 1 else "kappa_undefined_single_category" if n else "no_resolved_pairs",
    }


def bit(value, context):
    require(str(value) in ("0", "1"), f"{context}: expected 0/1, got {value!r}")
    return int(value)


def sampling(key):
    m, s, label = (key["module"], key["stratum"], key["label"])
    if m == "dossier":
        return (s, key["intent"], key["instrument_state"])
    if m == "intent":
        return (s, str(bit(label["own_trial_planned"], "intent instrument call")))
    return (s, key["instrument_state"])


def load_key(path, a):
    out, cells = ({}, defaultdict(list))
    for k in csv_rows(path, KEY_COLUMNS):
        idx = (k["module"], k["item"])
        require(idx in a and idx not in out, f"Unknown/duplicate key item {idx}")
        k["label"] = json.loads(k["instrument_label"])
        require(isinstance(k["label"], dict), f"{idx}: instrument_label must be an object")
        out[idx] = k
        m = k["module"]
        if m != "dossier_trial":
            require(
                k["stratum"]
                in (
                    ("pre_early", "pre_legacy", "post") if m in ("dossier", "intent") else ("pre_policy", "post_policy")
                ),
                f"{idx}: invalid stratum {k['stratum']!r}",
            )
            if m in ("dossier", "intent"):
                bit(k["intent"], f"{idx} intent")
            k["population"] = int(k["cell_population"])
            cells[m, sampling(k)].append(k)
    require(out.keys() == a.keys(), f"Key/workbook mismatch: {out.keys() ^ a.keys()}")
    for cell, members in cells.items():
        pops = {k["population"] for k in members}
        require(
            len(pops) == 1 and next(iter(pops)) >= len(members),
            f"{cell}: inconsistent/too-small cell_population {pops}, sampled={len(members)}",
        )
        for k in members:
            k["weight"] = k["population"] / len(members)
    return out


def observations(keys, ref, hu):
    lineage = indexed(csv_rows(hu / "lineage_labels.csv", ("item", "lineage")), "item", "lineage labels")
    obs = []
    for idx, k in keys.items():
        m, _ = idx
        r, lab = (ref[idx], k["label"])
        if m == "dossier_trial":
            continue
        if m == "dossier":
            new, kt = (bit(lab["primary"], "dossier primary"), bit(lab["ktrial_direct"], "dossier ktrial_direct"))
            state = "new_trial" if new else "ktrial_only" if kt else "none"
            require(k["instrument_state"] == state, f"{idx}: state/label mismatch")
            h = [None if r[f] == "unclear" else int(r[f] == "yes") for f in ("led_new_trial", "led_K_trial")]
            if r["search_done"] != "yes":
                h = [None, None]
            any_h = 1 if 1 in h else 0 if h == [0, 0] else None
            outcomes = [("new_trial", new, h[0]), ("K_trial", kt, h[1]), ("any_trial", int(new or kt), any_h)]
        elif m == "intent":
            outcomes = [
                (
                    "own_trial_planned",
                    bit(lab["own_trial_planned"], "intent"),
                    None if r["intent"] == "unclear" else int(r["intent"] == "own_trial_planned"),
                )
            ]
        elif m == "boundary":
            require(k["instrument_state"] in ("new", "grant_linked", "judged_only"), f"{idx}: bad boundary class")
            h = None if r["k_relationship"] == "unclear" else int(r["k_relationship"] == "new_protocol")
            pred = int(k["instrument_state"] == "new")
            outcomes = [("new_protocol", pred, h), ("original_K", 1 - pred, None if h is None else 1 - h)]
        else:
            bit(k["instrument_state"], "lineage has_ktrial sampling state")
            source = str(lab["source_item"])
            require(source in lineage, f"lineage source_item missing: {source}")
            pred = lineage[source]["lineage"]
            require(pred in CHOICES["lineage"], f"lineage invalid predicted lineage: {pred!r}")
            outcomes = [("developed_from_K", int(pred in DEV), int(r["lineage"] in DEV))]
        for outcome, pred, human in outcomes:
            obs.append(
                {
                    "module": m,
                    "outcome": outcome,
                    "stratum": k["stratum"],
                    "intent": k["intent"] if m in ("dossier", "intent") else "1",
                    "cell": sampling(k),
                    "weight": k["weight"],
                    "pred": pred,
                    "human": human,
                }
            )
    return obs


def metric(rows, name):
    resolved = [r for r in rows if r["human"] is not None]
    condition, success = {"Se": ("human", 1), "Sp": ("human", 0), "PPV": ("pred", 1), "NPV": ("pred", 0)}[name]
    den = [r for r in resolved if r[condition] == success]
    total = sum(r["weight"] for r in den)
    num = sum(r["weight"] for r in den if r["pred"] == r["human"])
    return (num / total if total else None, den, total)


def wilson(p, n):
    z = 1.959963984540054
    scale = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / scale
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / scale
    return (max(0, centre - half), min(1, centre + half))


def accuracy(obs):
    groups = defaultdict(list)
    for r in obs:
        for s, intent in (("overall", "all"), (r["stratum"], r["intent"])):
            groups[r["module"], r["outcome"], s, intent].append(r)
    out = []
    order = {name: i for i, name in enumerate(("dossier", "intent", "boundary", "lineage"))}
    for (module, outcome, s, intent), rows in sorted(
        groups.items(), key=lambda pair: (order[pair[0][0]], *pair[0][1:])
    ):
        for name in ("Se", "Sp", "PPV", "NPV"):
            p, den, weight = metric(rows, name)
            neff = weight**2 / sum(r["weight"] ** 2 for r in den) if den else 0
            lo, hi = wilson(p, neff) if den else ("", "")
            out.append(
                {
                    "module": module,
                    "outcome": outcome,
                    "stratum": s,
                    "intent": intent,
                    "metric": name,
                    "n_sampled": len(rows),
                    "n_resolved": sum(r["human"] is not None for r in rows),
                    "n_excluded": sum(r["human"] is None for r in rows),
                    "excluded_weight": sum(r["weight"] for r in rows if r["human"] is None),
                    "denominator_n": len(den),
                    "denominator_weight": weight,
                    "effective_n": neff,
                    "estimate": p if p is not None else "",
                    "ci_low": lo,
                    "ci_high": hi,
                    "ci_method": "wilson_kish_95",
                    "status": "ok" if den else "zero_denominator",
                }
            )
    return out


def write_csv(path, rows):
    require(rows, f"{path}: refusing empty output")
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print("aggregate rows:", len(rows))


def score(root):
    hu = DATA_DIR / "ratings"
    base = DATA_DIR / "validation"
    a, b, dis = ratings(base)
    ref = reference(base, a, dis)
    keys = load_key(base / "initial_key.csv", a)
    agreements = []
    for module, fields in FIELDS.items():
        for field in fields:
            agreements.append(
                agreement(module, field, [(r[field], b[idx][field]) for idx, r in a.items() if idx[0] == module])
            )
    obs = observations(keys, ref, hu)
    outputs = {"agreement.csv": agreements, "instrument_accuracy.csv": accuracy(obs)}
    out = RESULTS_DIR / "validation"
    out.mkdir(parents=True, exist_ok=True)
    for name, rows in outputs.items():
        write_csv(base / ("initial_" + name), rows)
        write_csv(out / name, [r for r in rows if r["module"] in ("intent", "boundary", "lineage")])
    return outputs


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=("merge", "score"))
    args = parser.parse_args()
    if args.mode == "merge":
        merge(DATA_DIR / "validation")
    else:
        score(ROOT)


if __name__ == "__main__":
    main()
    print("validation/score.py: complete")
