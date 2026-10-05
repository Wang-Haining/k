"Sample blinded dossier–lineage validation workbooks and write their private sampling key."

import csv
import hashlib
import json
import random
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config import DATA_DIR, DOSSIER_SAMPLE_SEED, ROOT
from outcomes.reporting import stratum

INSTRUMENT_DIR, HD, RV, HU = (
    DATA_DIR / "instrument",
    DATA_DIR / "historical",
    DATA_DIR / "derived",
    DATA_DIR / "ratings",
)
OUT = DATA_DIR / "validation"
RATERS = ("rater_a", "rater_b")
CAP = 30
CHOICES = {
    "identity": ["same_person", "different_person", "not_named", "unclear"],
    "role": [
        "overall_PI",
        "responsible_party_PI",
        "sponsor_investigator",
        "protocol_named_PI",
        "nonqualifying_role",
        "not_named",
        "unclear",
    ],
    "role_source": ["current_registry", "historical_snapshot", "protocol_excerpt", "none"],
    "scope": ["applied", "BESH_only", "no_human_intervention", "unclear"],
    "k_relationship": ["original_K", "new_protocol", "not_applicable", "unclear"],
    "start_in_window": ["yes", "no", "unclear"],
    "yes_no_unclear": ["yes", "no", "unclear"],
    "intent": ["own_trial_planned", "mentor_led_trial_explicit", "other_research", "unclear"],
    "lineage": ["extension", "same_approach_new_question", "related_topic", "unrelated"],
    "krel": ["original_K", "new_protocol", "unclear"],
}


def rows(p):
    with p.open() as handle:
        data = list(csv.DictReader(handle))
    assert data, "expected nonempty sampling input"
    return data


def officials(t):
    o = "; ".join(
        f"{x.get('name', '')} ({x.get('role', '')}, {x.get('affiliation', '')})"
        for x in t.get("overall_officials") or []
    )
    rp = t.get("responsible_party") or {}
    return (
        o,
        " ".join(
            str(rp.get(k, ""))
            for k in ("type", "investigatorFullName", "investigatorTitle", "investigatorAffiliation", "oldOrganization")
            if rp.get(k)
        ),
    )


def sample_dossiers(pa, extended, rng):
    ids = [b for b in pa if pa[b]["mechanism"] == "K23" and extended[b]["besh_only_classifier"] == "0"]

    def state(b):
        return (
            "new_trial" if pa[b]["primary"] == "1" else "ktrial_only" if extended[b]["ktrial_direct"] == "1" else "none"
        )

    cells = defaultdict(list)
    for b in sorted(ids):
        cells[stratum(pa[b]), pa[b]["intent"], state(b)].append(b)
    picks = []
    for c in sorted(cells):
        for b in rng.sample(cells[c], min(4, len(cells[c]))):
            picks.append((b, c, len(cells[c])))
    return (ids, picks)


def sample_intent(pa, ids, exclude, rng):
    qwen = {
        p.stem: json.loads(p.read_text())["label"]["research_type"] == "own_trial_planned"
        for p in (HD / "intent_labels").glob("*.json")
    }
    cells = defaultdict(list)
    for b in sorted(ids):
        if b not in exclude and b in qwen:
            cells[stratum(pa[b]), int(qwen[b])].append(b)
    alloc = {
        ("pre_early", 0): 17,
        ("pre_early", 1): 17,
        ("post", 0): 17,
        ("post", 1): 17,
        ("pre_legacy", 0): 16,
        ("pre_legacy", 1): 16,
    }
    return [
        (b, c, len(cells[c]), int(qwen[b]))
        for c in sorted(alloc)
        for b in rng.sample(cells[c], min(alloc[c], len(cells[c])))
    ]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rng = random.Random(DOSSIER_SAMPLE_SEED)
    pa = {r["blind_id"]: r for r in rows(HD / "person_analysis.csv")}
    extended = {r["blind_id"]: r for r in rows(RV / "historical_extended.csv")}
    persons = {r["blind_id"]: r for r in rows(HD / "persons.csv")}
    key = {
        k["blind_id"]: k
        for f in (DATA_DIR / "contemporaneous" / "cohort_key_private.json", HD / "cohort_key_private.json")
        for k in json.loads(f.read_text())
    }
    masked = {}
    for line in (HD / "intent_packets.jsonl").open():
        x = json.loads(line)
        masked[x["case_id"]] = x
    ids, dossier_sample = sample_dossiers(pa, extended, rng)
    dids = {b for b, _, _ in dossier_sample}
    intents = sample_intent(pa, ids, dids, rng)
    pairs = defaultdict(list)
    for r in rows(HD / "pairs.csv"):
        if r["blind_id"] in dids and r["window"] in ("inside", "boundary_uncertain", "missing"):
            pairs[r["blind_id"]].append(r)
    outc = {r["case_id"]: r for r in rows(HD / "pair_outcomes.csv") if r["blind_id"] in dids}
    shown, hidden = ({}, {})
    for b, ps in pairs.items():
        passed = [p for p in ps if p["qwen_reviewed"] == "True"]
        rest = [p for p in ps if p["qwen_reviewed"] != "True"]
        rng.shuffle(rest)
        keep = (passed + rest)[:CAP]
        rng.shuffle(keep)
        shown[b], hidden[b] = (keep, len(ps) - len(keep))
    need = {p["case_id"] for ps in shown.values() for p in ps}
    pk = {}
    for line in (HD / "packets.jsonl").open():
        cid = line[13 : 13 + 40].split('"', 1)[0]
        if cid in need:
            x = json.loads(line)
            pk[x["case_id"]] = x
    qlab = {}
    for c in need:
        f = HD / "leadership_labels" / f"{c}.json"
        if f.exists():
            qlab[c] = json.loads(f.read_text())["label"]

    def order(xs):
        return rng.sample(xs, len(xs))

    priv = []
    dossiers, dossier_trials = ([], [])
    for i, (b, cell, npop) in enumerate(order(dossier_sample), 1):
        item = f"D{i:02d}"
        k, m = (key[b], masked[b])
        dossiers.append(
            {
                "item": item,
                "awardee": f"{persons[b]['first_name']} {persons[b]['last_name']}",
                "K_institution": k["institution"],
                "K_project": k["core_project_num"],
                "K_start": k["start_date"][:10],
                "window_end_5y": k["end_5y"][:10],
                "K_title": m["K_title"],
                "K_abstract_masked": m["K_abstract"],
                "n_candidates_listed": len(shown.get(b, [])),
                "n_candidates_not_listed": hidden.get(b, 0),
                "ctgov_name_search": f"https://clinicaltrials.gov/search?term={persons[b]['first_name']}%20{persons[b]['last_name']}",
                "led_K_trial": "",
                "led_new_trial": "",
                "missed_trials_found": "",
                "search_done": "",
                "notes": "",
            }
        )
        priv.append(
            {
                "module": "dossier",
                "item": item,
                "blind_id": b,
                "case_id": "",
                "stratum": cell[0],
                "intent": cell[1],
                "instrument_state": cell[2],
                "cell_population": npop,
                "instrument_label": json.dumps(
                    {"primary": pa[b]["primary"], "ktrial_direct": extended[b]["ktrial_direct"]}
                ),
            }
        )
        for j, p in enumerate(shown.get(b, []), 1):
            t = pk[p["case_id"]]["trial"]
            o, rp = officials(t)
            dossier_trials.append(
                {
                    "item": item,
                    "trial": f"{item}.{j:02d}",
                    "nct_id": p["nct_id"],
                    "ctgov_url": f"https://clinicaltrials.gov/study/{p['nct_id']}",
                    "title": t.get("brief_title", ""),
                    "start_date": (t.get("start") or {}).get("date", ""),
                    "study_type": t.get("study_type", ""),
                    "overall_officials": o,
                    "responsible_party": rp,
                    "lead_sponsor": t.get("lead_sponsor", ""),
                    "secondary_ids": "; ".join(t.get("secondary_ids") or []),
                    "identity": "",
                    "role": "",
                    "scope": "",
                    "start_in_window": "",
                    "k_relationship": "",
                    "evidence": "",
                }
            )
            priv.append(
                {
                    "module": "dossier_trial",
                    "item": f"{item}.{j:02d}",
                    "blind_id": b,
                    "case_id": p["case_id"],
                    "stratum": cell[0],
                    "intent": cell[1],
                    "instrument_state": cell[2],
                    "cell_population": npop,
                    "instrument_label": json.dumps(
                        {
                            "screen_passed": p["qwen_reviewed"],
                            "window": p["window"],
                            "qwen": qlab.get(p["case_id"]),
                            "primary_pair": outc.get(p["case_id"], {}).get("primary_pair"),
                            "secondary_pair": outc.get(p["case_id"], {}).get("secondary_pair"),
                        }
                    ),
                }
            )
    intent = []
    for i, (b, cell, npop, call) in enumerate(order(intents), 1):
        item = f"I{i:03d}"
        intent.append(
            {
                "item": item,
                "K_title": masked[b]["K_title"],
                "K_abstract_masked": masked[b]["K_abstract"],
                "intent": "",
                "reason": "",
            }
        )
        priv.append(
            {
                "module": "intent",
                "item": item,
                "blind_id": b,
                "case_id": "",
                "stratum": cell[0],
                "intent": pa[b]["intent"],
                "instrument_state": "",
                "cell_population": npop,
                "instrument_label": json.dumps({"own_trial_planned": call}),
            }
        )
    ksheet = {r["item"]: r for r in rows(HU / "krel_rater_sheet.csv")}
    kcells = defaultdict(list)
    for r in rows(HU / "krel_sample_private.csv"):
        if r["analysis"] == "historical":
            kcells[r["period"], r["instrument_class"]].append(r)
    boundary = []
    picks = [(r, c, len(kcells[c])) for c in sorted(kcells) for r in rng.sample(kcells[c], min(10, len(kcells[c])))]
    for i, (r, c, npop) in enumerate(order(picks), 1):
        item = f"K{i:02d}"
        boundary.append(
            {
                "item": item,
                **{k: v for k, v in ksheet[r["item"]].items() if k not in ("item", "k_relationship", "reason")},
                "k_relationship": "",
                "reason": "",
            }
        )
        priv.append(
            {
                "module": "boundary",
                "item": item,
                "blind_id": "",
                "case_id": r["case_id"],
                "stratum": c[0],
                "intent": "",
                "instrument_state": c[1],
                "cell_population": npop,
                "instrument_label": json.dumps({"source_item": r["item"]}),
            }
        )
    lsheet = {r["item"]: r for r in rows(HU / "lineage_rater_sheet.csv")}
    lcells = defaultdict(list)
    for r in rows(HU / "lineage_sample_private.csv"):
        lcells[r["period"], r["has_ktrial"]].append(r)
    lalloc = {("pre_policy", "1"): 20, ("post_policy", "1"): 20, ("pre_policy", "0"): 10, ("post_policy", "0"): 10}
    picks = [
        (r, c, len(lcells[c])) for c in sorted(lalloc) for r in rng.sample(lcells[c], min(lalloc[c], len(lcells[c])))
    ]
    lineage = []
    for i, (r, c, npop) in enumerate(order(picks), 1):
        item = f"L{i:02d}"
        lineage.append(
            {
                "item": item,
                **{k: v for k, v in lsheet[r["item"]].items() if k not in ("item", "lineage", "reason")},
                "lineage": "",
                "reason": "",
            }
        )
        priv.append(
            {
                "module": "lineage",
                "item": item,
                "blind_id": r["blind_id"],
                "case_id": r["case_id"],
                "stratum": c[0],
                "intent": "",
                "instrument_state": c[1],
                "cell_population": npop,
                "instrument_label": json.dumps({"source_item": r["item"]}),
            }
        )
    pilot = []
    for rater in RATERS:
        write_book(OUT / f"rating_{rater}.xlsx", rater, pilot, dossiers, dossier_trials, intent, boundary, lineage)
    with (OUT / "key_private.csv").open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(priv[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(priv)
        print("CSV rows:", len(priv))
    manifest = {
        "seed": DOSSIER_SAMPLE_SEED,
        "cap_candidates": CAP,
        "pilot_items": pilot,
        "counts": {
            "dossier_awardees": len(dossiers),
            "dossier_trials": len(dossier_trials),
            "intent": len(intent),
            "boundary": len(boundary),
            "lineage": len(lineage),
        },
        "inputs_sha256": {
            str(p.relative_to(DATA_DIR)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in (
                HD / "person_analysis.csv",
                RV / "historical_extended.csv",
                HD / "pairs.csv",
                HU / "krel_rater_sheet.csv",
                HU / "lineage_rater_sheet.csv",
            )
        },
    }
    (OUT / "sample_manifest.json").write_text(json.dumps(manifest, indent=1))
    print("validation sample:", manifest["counts"])


ZH = ()


def md_plain(s):
    return re.sub("\\*\\*|`", "", s).strip()


def write_readme(ws, rater, pilot, banner=()):
    ws.title = "README"
    widths = {"A": 32, "B": 62, "C": 62, "D": 40}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    total = sum(widths.values())
    wrap = Alignment(wrap_text=True, vertical="top")

    def para(text, font=None, height=None):
        ws.append([text])
        r = ws.max_row
        ws.merge_cells(f"A{r}:D{r}")
        c = ws.cell(r, 1)
        c.alignment = wrap
        if font:
            c.font = font
        ws.row_dimensions[r].height = height or 15 * max(1, -(-len(text) // int(total * 0.95)))

    para(f"Rater / 评审: {rater}", Font(bold=True, size=14))
    for i, line in enumerate(ZH):
        para(line, Font(bold=True, size=12) if i == 0 else None)
    for line in banner:
        para(line, Font(bold=True, color="C00000", size=12))
    if pilot:
        para(f"Pilot items: {', '.join(pilot)}")
    para("")
    md = (ROOT / "validation" / "codebook.md").read_text()
    for line in md.splitlines():
        s = line.rstrip()
        if not s or s.startswith("\\"):
            continue
        if s.startswith("|"):
            cells = [md_plain(x) for x in s.strip("|").split("|")]
            if all(set(x) <= set("-: ") for x in cells):
                continue
            ws.append(cells[:4] if len(cells) <= 4 else cells[:3] + [" | ".join(cells[3:])])
            r = ws.max_row
            hdr = ws.cell(r - 1, 1).value is None or ws.cell(r - 1, 2).value is None
            for j, _x in enumerate(cells[:4], 1):
                c = ws.cell(r, j)
                c.alignment = wrap
                if hdr:
                    c.font = Font(bold=True)
            ws.row_dimensions[r].height = 15 * max(
                1, max((-(-len(x) // int(list(widths.values())[min(j, 3)] * 0.95)) for j, x in enumerate(cells[:4])))
            )
        elif s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            para(md_plain(s.lstrip("#")), Font(bold=True, size=14 if level == 1 else 12), 22)
        else:
            para(md_plain(s))


def write_book(path, rater, pilot, dossiers, dossier_trials, intent, boundary, lineage, banner=()):
    wb = Workbook()
    write_readme(wb.active, rater, pilot, banner)
    answer = {
        "led_K_trial": "yes_no_unclear",
        "led_new_trial": "yes_no_unclear",
        "search_done": "yes_no_unclear",
        "identity": "identity",
        "role": "role",
        "role_source": "role_source",
        "scope": "scope",
        "start_in_window": "start_in_window",
        "k_relationship": "k_relationship",
        "intent": "intent",
        "lineage": "lineage",
    }
    free = {"missed_trials_found", "notes", "evidence", "reason", "rater_notes"}
    sheets = (
        ("dossier_awardees", dossiers, {}),
        ("dossier_trials", dossier_trials, {}),
        ("intent", intent, {}),
        ("boundary", boundary, {"k_relationship": "krel"}),
        ("lineage", lineage, {}),
    )
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for name, data, override in sheets:
        ws = wb.create_sheet(name)
        cols = list(data[0])
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in data:
            ws.append([r.get(c, "") for c in cols])
        ws.freeze_panes = "C2"
        for j, c in enumerate(cols, 1):
            letter = ws.cell(1, j).column_letter
            long = c in (
                "K_abstract_masked",
                "K_abstract",
                "trial_summary",
                "Ktrial_summary",
                "new_summary",
                "overall_officials",
            )
            ws.column_dimensions[letter].width = 70 if long else 22
            kind = override.get(c, answer.get(c))
            if kind or c in free:
                for row in range(2, len(data) + 2):
                    ws.cell(row, j).fill = yellow
            if kind:
                dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES[kind]) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}{len(data) + 1}")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


if __name__ == "__main__":
    main()
    print("validation/sample.py: complete")
