"Pool adjudicated and assigned-rater references into agreement and accuracy estimates."

import json
from collections import Counter

from openpyxl import load_workbook

from config import DATA_DIR, RESULTS_DIR
from validation import score as E

BASE = DATA_DIR / "validation"
ART = RESULTS_DIR / "validation"
RATERS = ("rater_a", "rater_b")
AW = ("led_K_trial", "led_new_trial")
TR = ("identity", "role", "scope", "start_in_window", "k_relationship")


def additional(rater):
    wb = load_workbook(BASE / "additional_returns" / f"rating_{rater}.xlsx", read_only=True, data_only=True)
    out = {}
    for name, idc in (("dossier_awardees", "item"), ("dossier_trials", "trial")):
        rows = iter(wb[name].iter_rows(values_only=True))
        hdr = next(rows)
        for r in rows:
            d = {h: E.text(v) for h, v in zip(hdr, r, strict=False)}
            if int(d[idc].split(".")[0][1:]) > 70:
                out[name, d[idc]] = d
    wb.close()
    return out


def main():
    a, _, dis = E.ratings(BASE)
    ref1 = E.reference(BASE, a, dis)
    key = {
        r["item"]: r
        for r in E.csv_rows(
            BASE / "key_private.csv",
            ("module", "item", "stratum", "intent", "instrument_state", "cell_population", "instrument_label"),
        )
        if r["module"] == "dossier"
    }
    asg = {
        r["item"]: r for r in E.csv_rows(BASE / "assignment_private.csv", ("item", "set", "raters", "reference_rater"))
    }
    b2 = {r: additional(r) for r in RATERS}
    ov = [i for i, r in asg.items() if r["set"] == "overlap"]
    agr = []
    for f in AW:
        agr.append(
            E.agreement(
                "dossier_additional_overlap",
                f,
                [(b2["rater_a"]["dossier_awardees", i][f], b2["rater_b"]["dossier_awardees", i][f]) for i in ov],
            )
        )
    trials = [t for s, t in b2["rater_a"] if s == "dossier_trials" and t.split(".")[0] in set(ov)]
    for f in TR:
        agr.append(
            E.agreement(
                "dossier_trial_additional_overlap",
                f,
                [(b2["rater_a"]["dossier_trials", t][f], b2["rater_b"]["dossier_trials", t][f]) for t in trials],
            )
        )
    a1, b1, _ = E.ratings(BASE)
    w1 = [i for m, i in a1 if m == "dossier"]
    for f in AW:
        pairs = [(a1["dossier", i][f], b1["dossier", i][f]) for i in w1] + [
            (b2["rater_a"]["dossier_awardees", i][f], b2["rater_b"]["dossier_awardees", i][f]) for i in ov
        ]
        agr.append(E.agreement("dossier_all_double_rated", f, pairs))
    w1t = [i for m, i in a1 if m == "dossier_trial"]
    for f in TR:
        pairs = [(a1["dossier_trial", i][f], b1["dossier_trial", i][f]) for i in w1t] + [
            (b2["rater_a"]["dossier_trials", t][f], b2["rater_b"]["dossier_trials", t][f]) for t in trials
        ]
        agr.append(E.agreement("dossier_trial_all_double_rated", f, pairs))
    E.write_csv(BASE / "pooled_agreement_all.csv", agr)
    E.write_csv(ART / "pooled_agreement.csv", [r for r in agr if r["module"].endswith("_all_double_rated")])

    def reference():
        lab = {}
        for item in key:
            if int(item[1:]) <= 70:
                lab[item] = {f: ref1["dossier", item][f] for f in AW}
            else:
                r = asg[item]
                who = r["reference_rater"]
                lab[item] = {f: b2[who]["dossier_awardees", item][f] for f in AW}
        return lab

    n_cell = Counter((k["stratum"], k["intent"], k["instrument_state"]) for k in key.values())
    obs = []
    lab = reference()
    for item, k in key.items():
        hk, hp = (lab[item]["led_K_trial"], lab[item]["led_new_trial"])
        cell = (k["stratum"], k["intent"], k["instrument_state"])
        inst = json.loads(k["instrument_label"])
        w = float(k["cell_population"]) / n_cell[cell]
        pk, pp = (int(inst["ktrial_direct"]), int(inst["primary"]))
        for outcome, pred, h in (
            ("K_trial", pk, hk),
            ("new_trial", pp, hp),
            (
                "any_trial",
                int(pk or pp),
                "yes" if "yes" in (hk, hp) else "unclear" if "unclear" in (hk, hp) else "no",
            ),
        ):
            obs.append(
                {
                    "module": "dossier",
                    "outcome": outcome,
                    "stratum": k["stratum"],
                    "intent": k["intent"],
                    "weight": w,
                    "pred": pred,
                    "human": None if h == "unclear" else int(h == "yes"),
                }
            )
    E.write_csv(ART / "pooled_accuracy.csv", E.accuracy(obs))


if __name__ == "__main__":
    main()
    print("validation/pool.py: complete")
