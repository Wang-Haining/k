"Extend the private dossier sample to its specified cell sizes."

import csv
import json
import random
import shutil
from collections import defaultdict

from openpyxl import load_workbook

from config import DATA_DIR, DOSSIER_ADDITIONAL_SEED
from outcomes.reporting import stratum
from validation.sample import CAP, HD, OUT, RATERS, RV, officials, rows, write_book

TARGET = {
    ("post", "0", "ktrial_only"): 10,
    ("post", "0", "new_trial"): 19,
    ("post", "0", "none"): 22,
    ("post", "1", "ktrial_only"): 39,
    ("post", "1", "new_trial"): 33,
    ("post", "1", "none"): 8,
    ("pre_early", "0", "ktrial_only"): 10,
    ("pre_early", "0", "new_trial"): 15,
    ("pre_early", "0", "none"): 17,
    ("pre_early", "1", "ktrial_only"): 22,
    ("pre_early", "1", "new_trial"): 29,
    ("pre_early", "1", "none"): 8,
    ("pre_legacy", "0", "ktrial_only"): 2,
    ("pre_legacy", "0", "new_trial"): 10,
    ("pre_legacy", "0", "none"): 8,
    ("pre_legacy", "1", "ktrial_only"): 10,
    ("pre_legacy", "1", "new_trial"): 10,
    ("pre_legacy", "1", "none"): 8,
}
SHEETS = ("dossier_awardees", "dossier_trials", "intent", "boundary", "lineage")
BANNER = ("Complete the assigned additional dossiers; retain existing answers.",)


def sheet_rows(path):
    wb = load_workbook(path, data_only=True)
    out = {}
    for name in SHEETS:
        ws = wb[name]
        hdr = [c.value for c in ws[1]]
        out[name] = [
            {h: "" if v is None else v for h, v in zip(hdr, r, strict=False)}
            for r in ws.iter_rows(min_row=2, values_only=True)
        ]
    return out


def main():
    rng = random.Random(DOSSIER_ADDITIONAL_SEED)
    pa = {r["blind_id"]: r for r in rows(HD / "person_analysis.csv")}
    extended = {r["blind_id"]: r for r in rows(RV / "historical_extended.csv")}
    persons = {r["blind_id"]: r for r in rows(HD / "persons.csv")}
    key = {
        k["blind_id"]: k
        for f in (DATA_DIR / "contemporaneous" / "cohort_key_private.json", HD / "cohort_key_private.json")
        for k in json.loads(f.read_text())
    }
    if not (OUT / "initial_key.csv").exists():
        shutil.copy2(OUT / "key_private.csv", OUT / "initial_key.csv")
    priv = rows(OUT / "initial_key.csv")
    done = {r["blind_id"] for r in priv if r["module"] == "dossier"}
    ids = [b for b in pa if pa[b]["mechanism"] == "K23" and extended[b]["besh_only_classifier"] == "0"]

    def state(b):
        return (
            "new_trial" if pa[b]["primary"] == "1" else "ktrial_only" if extended[b]["ktrial_direct"] == "1" else "none"
        )

    cells = defaultdict(list)
    for b in sorted(ids):
        cells[stratum(pa[b]), pa[b]["intent"], state(b)].append(b)
    picks = []
    for c in sorted(TARGET):
        have = sum(b in done for b in cells[c])
        pool = [b for b in cells[c] if b not in done]
        picks += [(b, c, len(cells[c])) for b in rng.sample(pool, min(TARGET[c] - have, len(pool)))]
    rng.shuffle(picks)
    dids = {b for b, _, _ in picks}
    masked = {}
    for line in (HD / "intent_packets.jsonl").open():
        x = json.loads(line)
        if x["case_id"] in dids:
            masked[x["case_id"]] = x
    pairs = defaultdict(list)
    for r in rows(HD / "pairs.csv"):
        if r["blind_id"] in dids and r["window"] in ("inside", "boundary_uncertain", "missing"):
            pairs[r["blind_id"]].append(r)
    outc = {r["case_id"]: r for r in rows(HD / "pair_outcomes.csv") if r["blind_id"] in dids}
    shown, hidden = ({}, {})
    for b, ps in pairs.items():
        passed = [p for p in ps if p["qwen_reviewed"] == "True"]
        rest = [p for p in ps if p["qwen_reviewed"] != "True"]
        rng.shuffle(rest)
        keep = (passed + rest)[:CAP]
        rng.shuffle(keep)
        shown[b], hidden[b] = (keep, len(ps) - len(keep))
    need = {p["case_id"] for ps in shown.values() for p in ps}
    pk = {}
    for line in (HD / "packets.jsonl").open():
        cid = line[13:53].split('"', 1)[0]
        if cid in need:
            x = json.loads(line)
            pk[x["case_id"]] = x
    qlab = {
        c: json.loads((HD / "leadership_labels" / f"{c}.json").read_text())["label"]
        for c in need
        if (HD / "leadership_labels" / f"{c}.json").exists()
    }
    start = 1 + max(int(r["item"][1:]) for r in priv if r["module"] == "dossier")
    dossiers, dossier_trials = ([], [])
    for i, (b, cell, npop) in enumerate(picks, start):
        item = f"D{i:02d}"
        k, m = (key[b], masked[b])
        dossiers.append(
            {
                "item": item,
                "awardee": f"{persons[b]['first_name']} {persons[b]['last_name']}",
                "K_institution": k["institution"],
                "K_project": k["core_project_num"],
                "K_start": k["start_date"][:10],
                "window_end_5y": k["end_5y"][:10],
                "K_title": m["K_title"],
                "K_abstract_masked": m["K_abstract"],
                "n_candidates_listed": len(shown.get(b, [])),
                "n_candidates_not_listed": hidden.get(b, 0),
                "ctgov_name_search": f"https://clinicaltrials.gov/search?term={persons[b]['first_name']}%20{persons[b]['last_name']}",
                "led_K_trial": "",
                "led_new_trial": "",
                "missed_trials_found": "",
                "search_done": "",
                "notes": "",
            }
        )
        priv.append(
            {
                "module": "dossier",
                "item": item,
                "blind_id": b,
                "case_id": "",
                "stratum": cell[0],
                "intent": cell[1],
                "instrument_state": cell[2],
                "cell_population": npop,
                "instrument_label": json.dumps(
                    {"primary": pa[b]["primary"], "ktrial_direct": extended[b]["ktrial_direct"]}
                ),
            }
        )
        for j, p in enumerate(shown.get(b, []), 1):
            t = pk[p["case_id"]]["trial"]
            o, rp = officials(t)
            dossier_trials.append(
                {
                    "item": item,
                    "trial": f"{item}.{j:02d}",
                    "nct_id": p["nct_id"],
                    "ctgov_url": f"https://clinicaltrials.gov/study/{p['nct_id']}",
                    "title": t.get("brief_title", ""),
                    "start_date": (t.get("start") or {}).get("date", ""),
                    "study_type": t.get("study_type", ""),
                    "overall_officials": o,
                    "responsible_party": rp,
                    "lead_sponsor": t.get("lead_sponsor", ""),
                    "secondary_ids": "; ".join(t.get("secondary_ids") or []),
                    "identity": "",
                    "role": "",
                    "scope": "",
                    "start_in_window": "",
                    "k_relationship": "",
                    "evidence": "",
                }
            )
            priv.append(
                {
                    "module": "dossier_trial",
                    "item": f"{item}.{j:02d}",
                    "blind_id": b,
                    "case_id": p["case_id"],
                    "stratum": cell[0],
                    "intent": cell[1],
                    "instrument_state": cell[2],
                    "cell_population": npop,
                    "instrument_label": json.dumps(
                        {
                            "screen_passed": p["qwen_reviewed"],
                            "window": p["window"],
                            "qwen": qlab.get(p["case_id"]),
                            "primary_pair": outc.get(p["case_id"], {}).get("primary_pair"),
                            "secondary_pair": outc.get(p["case_id"], {}).get("secondary_pair"),
                        }
                    ),
                }
            )
    exp = OUT / "expand"
    exp.mkdir(exist_ok=True)
    for rater in RATERS:
        old = sheet_rows(OUT / "returned" / f"rating_{rater}.xlsx")
        write_book(
            exp / f"rating_{rater}.xlsx",
            rater,
            [],
            old["dossier_awardees"] + dossiers,
            old["dossier_trials"] + dossier_trials,
            old["intent"],
            old["boundary"],
            old["lineage"],
            BANNER,
        )
    with (OUT / "key_private.csv").open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(priv[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(priv)
        print("CSV rows:", len(priv))
    cnt = defaultdict(int)
    for r in priv:
        if r["module"] == "dossier":
            cnt[r["stratum"], r["intent"], r["instrument_state"]] += 1
    assert all(cnt[c] == min(TARGET[c], len(cells[c])) for c in TARGET), dict(cnt)


if __name__ == "__main__":
    main()
    print("validation/expand.py: complete")
